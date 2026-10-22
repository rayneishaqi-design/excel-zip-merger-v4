import csv
import io
import os
import shutil
import tempfile
import zipfile
from collections import Counter
from contextlib import closing
from pathlib import PurePosixPath

import streamlit as st
from openpyxl import load_workbook
import xlsxwriter

st.set_page_config(page_title='Penggabung Data Besar', page_icon='📦', layout='wide')
st.title('📦 Penggabung ZIP, Excel, dan CSV — Data Besar')
st.write('Gabungkan banyak ZIP/Excel/CSV dengan pemrosesan bertahap. Tidak ada batas jumlah file/baris buatan aplikasi, tetapi server tetap memiliki batas CPU, RAM, disk dan upload.')
st.warning('Jangan unggah data perusahaan rahasia ke layanan publik tanpa izin. File diproses pada server aplikasi.')
st.caption('Hasil CSV dipotong menjadi bagian maksimal 700.000 baris agar mudah dibuka di Excel. Hasil XLSX otomatis dibagi ke beberapa sheet (maksimal 1.000.000 baris data per sheet).')

MAX_CSV_ROWS = 700_000
MAX_XLSX_ROWS = 1_000_000
EXTS = ('.csv', '.xlsx', '.xlsm')


def clean_header(vals):
    result = []
    seen = Counter()
    for i, value in enumerate(vals, 1):
        head = str(value).strip() if value is not None else ''
        head = head or f'Kolom_{i}'
        seen[head] += 1
        if seen[head] > 1:
            head = f'{head}__{seen[head]}'
        result.append(head)
    return result


def find_header(row):
    return any(v is not None and str(v).strip() for v in row)


def text_cell(v):
    if v is None:
        return ''
    if hasattr(v, 'isoformat'):
        return v.isoformat(sep=' ') if hasattr(v, 'hour') else v.isoformat()
    return str(v)


def keep_safe_zip(info):
    name = info.filename.replace('\\', '/')
    parts = PurePosixPath(name).parts
    if info.is_dir() or '..' in parts or name.startswith('/') or '__MACOSX' in parts:
        return False
    return name.lower().endswith(EXTS)


def copy_to_spool(source, threshold=16 * 1024 * 1024):
    f = tempfile.SpooledTemporaryFile(max_size=threshold, mode='w+b')
    shutil.copyfileobj(source, f, length=1024 * 1024)
    f.seek(0)
    return f


def iter_upload_members(uploads):
    for uploaded in uploads:
        uploaded.seek(0)
        if uploaded.name.lower().endswith('.zip'):
            with zipfile.ZipFile(uploaded) as outer:
                for info in outer.infolist():
                    if not keep_safe_zip(info):
                        continue
                    if info.flag_bits & 1:
                        raise ValueError(f'ZIP berpassword: {info.filename}')
                    with outer.open(info) as member:
                        with copy_to_spool(member) as f:
                            yield f'{uploaded.name}/{info.filename}', f
        elif uploaded.name.lower().endswith(EXTS):
            yield uploaded.name, uploaded


def csv_rows(source):
    source.seek(0)
    head = source.read(4096)
    source.seek(0)
    enc = 'utf-8-sig'
    for name in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            head.decode(name, errors='strict')
            enc = name
            break
        except UnicodeDecodeError:
            pass
    wrapper = io.TextIOWrapper(source, encoding=enc, errors='replace', newline='')
    try:
        sample = wrapper.read(16_384)
        wrapper.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        yield from csv.reader(wrapper, dialect)
    finally:
        try:
            wrapper.detach()
        except (ValueError, OSError):
            pass


def iterate_sheets(file_name, source, mode, wanted_sheet):
    if file_name.lower().endswith('.csv'):
        yield 'CSV', csv_rows(source)
        return
    source.seek(0)
    with closing(load_workbook(source, read_only=True, data_only=True)) as wb:
        if mode == 'Sheet pertama':
            sheets = wb.worksheets[:1]
        elif mode == 'Nama sheet tertentu':
            sheets = [s for s in wb.worksheets if s.title.lower() == wanted_sheet.lower()]
        else:
            sheets = wb.worksheets
        for ws in sheets:
            # Penting: beberapa file sumber menyimpan dimensi 'A1:A1' yang salah.
            # reset_dimensions memaksa openpyxl membaca baris asli.
            ws.reset_dimensions()
            yield ws.title, ws.iter_rows(values_only=True)


def scan_columns(uploads, mode, wanted_sheet, metadata):
    headers = []
    seen = set()
    checked = 0
    problems = []
    for path, handle in iter_upload_members(uploads):
        try:
            for sheet, rows in iterate_sheets(path, handle, mode, wanted_sheet):
                for first in rows:
                    if find_header(first):
                        for col in clean_header(first):
                            if col not in seen:
                                seen.add(col)
                                headers.append(col)
                        checked += 1
                        break
                if hasattr(rows, 'close'):
                    rows.close()
        except Exception as exc:
            problems.append([path, '', f'Gagal pemeriksaan header: {type(exc).__name__}: {exc}'])
    if metadata:
        for field in ('__asal_file', '__asal_sheet'):
            if field in seen:
                raise ValueError(f'Kolom sumber bernama {field} sudah ada, matikan metadata atau ubah nama kolom sumber.')
        headers = ['__asal_file', '__asal_sheet'] + headers
    return headers, checked, problems


class OutputWriter:
    def __init__(self, fmt, out_folder, headers):
        self.fmt = fmt
        self.folder = out_folder
        self.headers = headers
        self.count = 0
        self.part = 0
        self.paths = []
        self.wb = None
        self.worksheet = None
        self.file = None
        self.writer = None
        self.local_row = 0
        if fmt == 'Excel (.xlsx)':
            self.path = os.path.join(out_folder, 'hasil_gabungan.xlsx')
            self.wb = xlsxwriter.Workbook(self.path, {'constant_memory': True, 'strings_to_formulas': False, 'strings_to_urls': False})
        else:
            self.path = os.path.join(out_folder, 'hasil_csv_terbagi.zip')
            self.zf = zipfile.ZipFile(self.path, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=1, allowZip64=True)
        self.next_part()

    def next_part(self):
        if self.fmt == 'Excel (.xlsx)':
            self.part += 1
            self.worksheet = self.wb.add_worksheet(f'Data_{self.part}')
            self.worksheet.write_row(0, 0, self.headers)
        else:
            if self.file is not None:
                self.file.close()
            self.part += 1
            self.file = self.zf.open(f'gabungan_part_{self.part:03}.csv', 'w', force_zip64=True)
            self.text = io.TextIOWrapper(self.file, encoding='utf-8-sig', newline='', write_through=True)
            self.writer = csv.writer(self.text)
            self.writer.writerow(self.headers)
        self.local_row = 0

    def add(self, values):
        cap = MAX_XLSX_ROWS if self.fmt == 'Excel (.xlsx)' else MAX_CSV_ROWS
        if self.local_row >= cap:
            if self.fmt != 'Excel (.xlsx)':
                self.text.flush()
                self.text.detach()
            self.next_part()
        if self.fmt == 'Excel (.xlsx)':
            for i, v in enumerate(values):
                self.worksheet.write_string(self.local_row + 1, i, text_cell(v)[:32767])
        else:
            self.writer.writerow([text_cell(v) for v in values])
        self.count += 1
        self.local_row += 1

    def close(self):
        if self.fmt == 'Excel (.xlsx)':
            self.wb.close()
        else:
            self.text.flush()
            self.text.detach()
            self.file.close()
            self.zf.close()


def merge(uploads, mode, wanted_sheet, metadata, fmt, progress):
    headers, inspected, errors = scan_columns(uploads, mode, wanted_sheet, metadata)
    if not headers or (metadata and len(headers) == 2) or inspected == 0:
        raise ValueError('Tidak ada header/sheet yang terbaca. Periksa isi ZIP dan pilihan sheet.')
    folder = tempfile.mkdtemp(prefix='merger_large_')
    output = OutputWriter(fmt, folder, headers)
    total = 0
    report = []
    try:
        for number, (path, source) in enumerate(iter_upload_members(uploads), 1):
            progress.caption(f'Memproses file ke-{number}: {path[-90:]}')
            try:
                found_sheet = False
                for sheet, rows in iterate_sheets(path, source, mode, wanted_sheet):
                    found_sheet = True
                    count = 0
                    it = iter(rows)
                    first = next((row for row in it if find_header(row)), None)
                    if first is None:
                        report.append([path, sheet, 0, 'Sheet kosong'])
                        continue
                    current_headers = clean_header(first)
                    mapper = {col: i for i, col in enumerate(current_headers)}
                    fields = headers[2:] if metadata else headers
                    idx = [mapper.get(col) for col in fields]
                    for row in it:
                        if not find_header(row):
                            continue
                        values = [row[j] if j is not None and j < len(row) else '' for j in idx]
                        if metadata:
                            values = [path, sheet] + values
                        output.add(values)
                        count += 1
                    total += count
                    report.append([path, sheet, count, 'OK'])
                if not found_sheet:
                    errors.append([path, '', 'Tidak ada nama sheet yang cocok'])
            except Exception as exc:
                errors.append([path, '', f'{type(exc).__name__}: {exc}'])
        output.close()
        report_path = os.path.join(folder, 'laporan.csv')
        with open(report_path, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f)
            w.writerow(['Asal file', 'Asal sheet', 'Baris tergabung', 'Status'])
            w.writerows(report)
            w.writerow(['TOTAL', '', total, ''])
            w.writerow([])
            w.writerow(['FILE BERMASALAH', 'SHEET', 'KETERANGAN'])
            w.writerows(errors)
        return {'folder': folder, 'output': output.path, 'report': report_path, 'total': total, 'parts': output.part, 'nerrors': len(errors), 'nfiles': len(report), 'headers': len(headers)}
    except Exception:
        try:
            output.close()
        except Exception:
            pass
        shutil.rmtree(folder, ignore_errors=True)
        raise


files = st.file_uploader('Pilih beberapa ZIP / Excel / CSV (bisa dicampur)', type=['zip', 'xlsx', 'xlsm', 'csv'], accept_multiple_files=True)
col1, col2 = st.columns(2)
with col1:
    mode = st.selectbox('Sheet yang diproses', ['Semua sheet', 'Sheet pertama', 'Nama sheet tertentu'])
with col2:
    sheet = st.text_input('Nama sheet', 'orders', disabled=mode != 'Nama sheet tertentu')
metadata = st.checkbox('Tambahkan asal file dan sheet', value=True)
fmt = st.radio('Format hasil', ['CSV terpecah dalam ZIP (disarankan untuk jutaan baris)', 'Excel (.xlsx)'], horizontal=True)

if st.button('Gabungkan semua file', type='primary', disabled=not files):
    if 'last' in st.session_state:
        shutil.rmtree(st.session_state.last['folder'], ignore_errors=True)
        del st.session_state.last
    state = st.empty()
    try:
        with st.spinner('Memproses file secara bertahap...'):
            choice = 'Excel (.xlsx)' if fmt == 'Excel (.xlsx)' else 'CSV'
            result = merge(files, mode, sheet, metadata, choice, state)
        st.session_state.last = result
    except Exception as exc:
        st.error(f'Proses gagal: {type(exc).__name__}: {exc}')

if 'last' in st.session_state:
    data = st.session_state.last
    a, b, c, d = st.columns(4)
    a.metric('Baris gabungan', f"{data['total']:,}")
    b.metric('Sheet sumber terbaca', data['nfiles'])
    c.metric('Bagian hasil', data['parts'])
    d.metric('Masalah', data['nerrors'])
    if data['total'] == 0:
        st.error('Tidak ada baris data yang terbaca! Jangan gunakan output sebagai hasil yang valid. Periksa laporan dan sheet input.')
    else:
        st.success('Proses selesai. Bandingkan total baris dan file sumber melalui laporan sebelum memakai hasil.')
        with open(data['output'], 'rb') as f:
            st.download_button('📥 Download gabungan', f, file_name=os.path.basename(data['output']), mime='application/zip' if data['output'].endswith('.zip') else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    with open(data['report'], 'rb') as f:
        st.download_button('📋 Download laporan baris per file', f, file_name='laporan_penggabungan.csv', mime='text/csv')
    st.caption('Jika file hasil CSV dibuka langsung dan seluruh isi muncul di satu kolom, impor melalui Excel → Data → From Text/CSV → delimiter Comma. Nomor pesanan/barcode sebaiknya diimpor sebagai Text.')
